# -*- coding: utf-8 -*-
"""Экспорт скачанных реестров Минюста в registries.xlsx — каждый реестр на отдельной вкладке."""
import json
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TZ = timezone.utc


def fmt_date(v):
    """epoch ms → дд.мм.гггг"""
    if v in (None, "", 0):
        return ""
    try:
        ms = int(v)
    except (TypeError, ValueError):
        return str(v)
    if ms < 10**12:  # уже секунды? нет в этих данных, но на всякий случай
        ms *= 1000
    return datetime.fromtimestamp(ms / 1000, TZ).astimezone().strftime("%d.%m.%Y")


def load_cols(registry_id):
    data = json.load(open("registries/all_post.json", encoding="utf-8"))
    for v in data["values"]:
        if v["id"] == registry_id:
            return [(c["name"], c["title"], c["type"]) for c in v.get("columns", [])]
    return []


def write_sheet(ws, cols, rows, date_fields, width_map):
    headers = ["№ строки"] + [t for _, t, _ in cols if _ != "field_1_i"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for i, r in enumerate(rows, 1):
        line = [i]
        for name, _, typ in cols:
            v = r.get(name)
            if typ == "date":
                v = fmt_date(v)
            elif v is None:
                v = ""
            elif not isinstance(v, str):
                v = str(v)
            line.append(v.strip() if isinstance(v, str) else v)
        ws.append(line)

    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = width_map.get(i, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def main():
    wb = Workbook()
    wb.remove(wb.active)

    WIDE = 55

    # 1. Иностранные агенты
    fa = json.load(open("registries/foreign_agents.json", encoding="utf-8"))
    cols = load_cols("39b95df9-9a68-6b6d-e1e3-e6388507067e")
    ws = wb.create_sheet("Иностранные агенты")
    date_fields = {n for n, _, t in cols if t == "date"}
    write_sheet(ws, cols, fa["rows"], date_fields,
                {1: 7, 2: WIDE + 20, 3: 30, 4: 14, 5: 14, 6: 22, 7: 28, 8: 14, 12: 12, 13: WIDE, 14: 35})

    # 2. Экстремистские организации
    eo = json.load(open("registries/extremist_orgs.json", encoding="utf-8"))
    cols = load_cols("59961ebb-9f0e-d885-7ebc-b90399923b03")
    ws = wb.create_sheet("Экстремистские организации")
    write_sheet(ws, cols, eo["rows"], set(),
                {1: 8, 2: WIDE + 20, 3: 25, 4: 18, 5: 45, 6: 15, 7: 15})

    # 3. Нежелательные организации
    un = json.load(open("registries/undesirable_orgs.json", encoding="utf-8"))
    cols = load_cols("c2d1692e-a9f6-5a79-13ee-5da5b42980df")
    ws = wb.create_sheet("Нежелательные организации")
    write_sheet(ws, cols, un["rows"], set(),
                {1: 7, 2: 14, 3: 14, 4: 14, 5: WIDE + 20, 6: 15, 7: 14, 8: 14, 9: 14, 10: 14})

    # 4. Экстремистские материалы
    em = json.load(open("registries/extremist_materials.json", encoding="utf-8"))
    ws = wb.create_sheet("Экстремистские материалы")
    ws.append(["№ п/п", "Материал", "Дата"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in em["rows"]:
        ws.append([r["num"], r["material"], r["date"]])
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 120
    ws.column_dimensions["C"].width = 14
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"

    out = "registries.xlsx"
    wb.save(out)
    print(f"saved {out}: иноагенты={len(fa['rows'])}, экстр.орг={len(eo['rows'])}, "
          f"нежелат.орг={len(un['rows'])}, материалы={len(em['rows'])}")


if __name__ == "__main__":
    main()
